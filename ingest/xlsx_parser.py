"""
XLSX parser using openpyxl.
Extraction type: DETERMINISTIC — direct cell value reading from the XML-based
xlsx format. Handles merged cells by propagating the merge-source value.
"""
import openpyxl
from schema import (
    RawExtraction, ContentBlock, TableData, ExtractionStep,
    ExtractionMethod, MethodType
)


def parse_xlsx(filepath: str, file_info: dict) -> RawExtraction:
    """Parse an XLSX file using openpyxl."""
    extraction = RawExtraction(
        source_file=file_info["path"],
        file_format="xlsx",
        trust_tier=file_info["trust_tier"],
        trust_label=file_info["trust_label"],
    )
    
    extraction.extraction_steps.append(ExtractionStep(
        step_name="xlsx_parsing",
        method=ExtractionMethod.OPENPYXL_XLSX,
        method_type=MethodType.DETERMINISTIC,
        reason="openpyxl reads cell values directly from the XLSX XML structure. "
               "No ML or heuristics — deterministic read of cell data. "
               "Merged cell handling uses openpyxl's merge range API to propagate values."
    ))
    
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        extraction.metadata["sheet_names"] = wb.sheetnames
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # Track merged cells
            merged_ranges = list(ws.merged_cells.ranges)
            has_merged = len(merged_ranges) > 0
            if has_merged:
                extraction.metadata[f"merged_cells_{sheet_name}"] = [str(r) for r in merged_ranges]
            
            # Build a merged-cell lookup: for any cell in a merged range,
            # return the value of the top-left cell of that range
            merge_lookup = {}
            for merge_range in merged_ranges:
                min_row, min_col = merge_range.min_row, merge_range.min_col
                top_left_val = ws.cell(row=min_row, column=min_col).value
                for row in range(merge_range.min_row, merge_range.max_row + 1):
                    for col in range(merge_range.min_col, merge_range.max_col + 1):
                        merge_lookup[(row, col)] = top_left_val
            
            # Read all rows, resolving merged cells
            all_rows = []
            for row_idx, row in enumerate(ws.iter_rows(values_only=False), 1):
                vals = []
                for col_idx, cell in enumerate(row, 1):
                    if cell.value is not None:
                        vals.append(str(cell.value).strip())
                    elif (row_idx, col_idx) in merge_lookup:
                        # Merged cell — use the source value
                        val = merge_lookup[(row_idx, col_idx)]
                        vals.append(str(val).strip() if val else "")
                    else:
                        vals.append("")
                all_rows.append(vals)
            
            # Skip fully empty rows
            all_rows = [r for r in all_rows if any(c.strip() for c in r)]
            
            if all_rows:
                headers = all_rows[0]
                data_rows = all_rows[1:]
                
                extraction.tables.append(TableData(
                    headers=headers,
                    rows=data_rows,
                    section=sheet_name,
                    has_merged_cells=has_merged,
                ))
                
                # Also add the sheet content as a content block for text search
                extraction.content_blocks.append(ContentBlock(
                    text=f"Sheet: {sheet_name} | Headers: {', '.join(headers)} | "
                         f"Rows: {len(data_rows)}",
                    block_type="table_summary",
                    section=sheet_name,
                ))
        
        wb.close()
        extraction.overall_confidence = 0.95
    
    except Exception as e:
        extraction.failures.append(f"XLSX parsing failed: {str(e)}")
        extraction.overall_confidence = 0.0
    
    return extraction
